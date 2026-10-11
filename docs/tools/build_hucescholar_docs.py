from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs' / 'deliverables'
QA = ROOT / 'docs' / '_qa'
OUT.mkdir(parents=True, exist_ok=True)
QA.mkdir(parents=True, exist_ok=True)
DATA = json.loads((Path(__file__).parent / 'hucescholar_content.json').read_text(encoding='utf-8'))

# Resolve repeated matrices from one source to keep the two documents consistent.
for section in DATA['proposal_sections']:
    for block in section['blocks']:
        if 'table' in block and not block['table']['rows']:
            block['table']['rows'] = DATA['assumptions'] if section['title'].startswith('3 ') else DATA['requirements']
for section in DATA['usecase_front']:
    for block in section['blocks']:
        if 'table' in block and not block['table']['rows']:
            block['table']['rows'] = DATA['assumptions'] if section['title'].startswith('1 ') else [
                [x['id'], x['name'], x['actors'].split(';')[0], ', '.join(x['fr'])] for x in DATA['usecases']
            ]
for section in DATA['usecase_back']:
    for block in section['blocks']:
        if 'table' in block and not block['table']['rows']:
            block['table']['rows'] = [
                [row[0], row[1], ', '.join(x['id'] for x in DATA['usecases'] if row[0] in x['fr'])]
                for row in DATA['requirements']
            ]


def field(p, instr):
    run = p.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instr
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    result = OxmlElement('w:t'); result.text = ' '
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for x in [begin, text, sep, result, end]: run._r.append(x)


def new_doc(kind):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21)
    sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.1); sec.right_margin = Cm(2.1)
    sec.header_distance = Cm(.8); sec.footer_distance = Cm(.8)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number', 'Caption']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Times New Roman')
    normal = doc.styles['Normal']
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    # The bundled Word base contains a theme border on Title. Remove it explicitly.
    for style in doc.styles:
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number', 'Caption']:
        fonts = doc.styles[name].element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn('w:' + attr), None)
        for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + attr), 'Times New Roman')
    for name, size, before, after in [('Title', 26, 0, 14), ('Subtitle', 14, 0, 10), ('Heading 1', 16, 14, 9), ('Heading 2', 13, 10, 6), ('Heading 3', 12, 8, 4)]:
        s = doc.styles[name]; s.font.size = Pt(size)
        s.font.bold = name != 'Subtitle'
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for name in ['List Bullet', 'List Number']:
        doc.styles[name].paragraph_format.space_after = Pt(3)
    header = sec.header.paragraphs[0]
    header.text = f'HUCEScholar  |  {kind}'
    header.runs[0].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('Phiên bản 1.1  •  09/10/2026  |  Trang '); r.font.size = Pt(9)
    field(footer, ' PAGE ')
    doc.core_properties.title = f'HUCEScholar {kind}'
    doc.core_properties.subject = 'Ứng dụng giáo trình điện tử dành cho sinh viên Đại học Xây dựng Hà Nội'
    doc.core_properties.author = 'HUCEScholar'
    doc.core_properties.keywords = 'HUCEScholar, proposal, use case, giáo trình điện tử'
    settings = doc.settings.element
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
    return doc


def p(doc, text, style=None):
    return doc.add_paragraph(text, style)


def bullets(doc, items):
    for x in items: p(doc, x, 'List Bullet')


def heading(doc, title, level=1):
    return doc.add_heading(title, level)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None: widths = [16.8 / len(headers)] * len(headers)
    for col, width in zip(t.columns, widths): col.width = Cm(width)
    for cell, text, width in zip(t.rows[0].cells, headers, widths): cell.text = text; cell.width = Cm(width)
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, text, width in zip(cells, row, widths): cell.text = str(text); cell.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement(f'w:{side}'); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '5'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for idx, row in enumerate(t.rows):
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, val in [('top', '90'), ('bottom', '90'), ('left', '100'), ('right', '100')]:
                el = OxmlElement(f'w:{side}'); el.set(qn('w:w'), val); el.set(qn('w:type'), 'dxa'); margins.append(el)
            props.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '233B53' if idx == 0 else ('F3F6F9' if idx % 2 == 0 else 'FFFFFF')); props.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.line_spacing = 1.05
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.space_before = Pt(1)
                if idx == 0 or (len(rows) <= 6 and idx < len(t.rows) - 1):
                    par.paragraph_format.keep_with_next = True
                if widths[col_idx] <= 2.2: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = Pt(10.5)
                    run.font.name = 'Times New Roman'
                    if idx == 0: run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    p(doc, '').paragraph_format.space_after = Pt(2)
    return t


def hyperlink(doc, label, url):
    par = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    rel = par.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79'); props.append(color)
    run.append(props); text = OxmlElement('w:t'); text.text = label; run.append(text); link.append(run); par._p.append(link)


def cover(doc, title, subtitle, intro):
    p(doc, title, 'Title')
    p(doc, subtitle, 'Subtitle')
    p(doc, 'Phiên bản 1.1  |  Ngày 09/10/2026  |  Ngôn ngữ Tiếng Việt')
    p(doc, 'Trạng thái Chính sách MVP đã chốt phục vụ thiết kế')
    for x in intro: p(doc, x)
    heading(doc, 'Cơ sở lập tài liệu', 2)
    p(doc, 'Tên HUCEScholar, định hướng giáo trình điện tử và ba vai trò ADMIN, LECTURER, STUDENT đã được người đề xuất xác nhận. Cơ sở nội bộ là trao đổi với người đề xuất và bản khởi động HUCEScholar phiên bản 0.4. A01–A06 được người đề xuất chốt ngày 09/10/2026; BR16–BR17 cụ thể hóa quyền admin và bàn giao. Quyết định phạm vi dự án không phải cam kết hợp tác hoặc quyền sử dụng dữ liệu của nhà trường.')
    doc.add_page_break()
    p(doc, 'Mục lục', 'Title')
    field(p(doc, ''), ' TOC \\o "1-1" \\h \\z \\u ')
    p(doc, 'Các heading và mã use case được giữ để tra cứu trong Navigation Pane của Word.')
    doc.add_page_break()


def section_content(doc, sections):
    for section in sections:
        if section.get('break_before'):
            doc.add_page_break()
        heading(doc, section['title'])
        for block in section['blocks']:
            if 'h2' in block: heading(doc, block['h2'], 2)
            if 'p' in block: p(doc, block['p'])
            if 'bullets' in block: bullets(doc, block['bullets'])
            if 'table' in block:
                spec = block['table']; table(doc, spec['headers'], spec['rows'], spec.get('widths'))
            if 'links' in block:
                for label, url in block['links']: hyperlink(doc, label, url)


def draw_diagram(actor, ids, filename):
    # UML association overview. Each image is a view of the same system boundary.
    w, h = 1800, 1600
    im = Image.new('RGB', (w,h), 'white'); d = ImageDraw.Draw(im)
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 32)
    small = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 28)
    bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 36)
    d.rounded_rectangle((420,90,1740,1520), radius=12, outline='#233B53', width=4)
    d.text((470,115), 'HUCEScholar', font=bold, fill='black')
    ax, ay = 210, 760
    d.ellipse((ax-38, ay-140, ax+38, ay-64), outline='black', width=4)
    d.line((ax,ay-64,ax,ay+65), fill='black', width=4)
    d.line((ax-65,ay-12,ax+65,ay-12), fill='black', width=4)
    d.line((ax,ay+65,ax-55,ay+145), fill='black', width=4)
    d.line((ax,ay+65,ax+55,ay+145), fill='black', width=4)
    label_box = d.textbbox((0,0), actor, font=bold)
    d.text((ax-(label_box[2]-label_box[0])/2, ay+175), actor, font=bold, fill='black')
    by_id = {x['id']: x for x in DATA['usecases']}
    top, bottom = 260, 1400
    for idx, uid in enumerate(ids):
        cy = top + idx * ((bottom-top)/max(1,len(ids)-1))
        left, right = 600, 1640
        d.line((ax+68,ay-12,left,cy), fill='#637487', width=3)
        d.ellipse((left,cy-62,right,cy+62), fill='#F3F6F9', outline='#233B53', width=3)
        label = f"{uid}  {by_id[uid]['short']}"
        bbox = d.textbbox((0,0), label, font=font)
        d.text(((left+right-(bbox[2]-bbox[0]))/2,cy-20), label, font=font, fill='black')
    im.save(filename)


def proposal():
    doc = new_doc('Project Proposal')
    cover(doc, 'HUCEScholar Project Proposal', 'Đề xuất dự án ứng dụng giáo trình điện tử', DATA['proposal_intro'])
    section_content(doc, DATA['proposal_sections'])
    path = OUT / 'HUCEScholar_Project_Proposal.docx'; doc.save(path)
    return path


def usecases():
    doc = new_doc('Use Case Specification')
    cover(doc, 'HUCEScholar Use Case Specification', 'Đặc tả use case theo ba vai trò', DATA['usecase_intro'])
    section_content(doc, DATA['usecase_front'])
    doc.add_page_break()
    heading(doc, '4 Sơ đồ use case theo vai trò')
    p(doc, 'Ba hình là các góc nhìn của cùng hệ thống. Đường liền thể hiện tác nhân tham gia use case, không thể hiện trình tự hay luồng dữ liệu. UC01, UC02 và UC03 dùng chung cho ba vai trò và được nêu trong danh mục; các quyền đọc của giảng viên/quản trị viên vẫn chịu chính sách tài liệu. Các hình tập trung vào trách nhiệm chính để tránh sơ đồ có quá nhiều đường nối.')
    views = [
        ('Quản trị viên', ['UC04','UC05','UC06','UC18','UC19','UC21','UC22']),
        ('Giảng viên', ['UC07','UC08','UC12','UC15','UC16','UC17']),
        ('Sinh viên', ['UC07','UC08','UC09','UC10','UC11','UC12','UC13','UC14','UC20']),
    ]
    for i,(actor, ids) in enumerate(views):
        if i > 0: doc.add_page_break()
        heading(doc, f'4 {i+1} Use case của {actor.lower()}', 2)
        img = QA / f'usecase_role_{i+1}.png'; draw_diagram(actor, ids, img)
        par = p(doc, ''); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(); run.add_picture(str(img), width=Cm(15.7))
        drawing = run._r.xpath('.//wp:docPr')[0]
        drawing.set('descr', f'Sơ đồ tác nhân {actor} liên kết tới ' + ', '.join(ids) + ' trong ranh giới HUCEScholar')
        p(doc, f'Hình {i+1} Các use case chính của {actor.lower()}', 'Caption')
    doc.add_page_break()
    heading(doc, '5 Đặc tả use case chi tiết')
    p(doc, 'Mỗi use case mô tả mục tiêu người dùng, tác nhân, điều kiện, luồng thành công, nhánh thay thế và kết quả. Phân quyền luôn được kiểm tra tại máy chủ; một màn hình bị ẩn không thay thế việc kiểm tra quyền.')
    for uc in DATA['usecases']:
        heading(doc, f"{uc['id']} {uc['name']}", 2)
        p(doc, f"Mục tiêu: {uc['goal']}")
        p(doc, f"Tác nhân: {uc['actors']}  |  Yêu cầu: {', '.join(uc['fr'])}")
        p(doc, f"Tiền điều kiện: {uc['pre']}")
        p(doc, f"Kích hoạt: {uc['trigger']}")
        label = p(doc, 'Luồng chính'); label.runs[0].bold = True; label.paragraph_format.keep_with_next = True
        for i, step in enumerate(uc['steps'], 1):
            step_p = p(doc, f'{i}. {step}')
            step_p.paragraph_format.left_indent = Cm(.35)
            step_p.paragraph_format.space_after = Pt(3)
        label = p(doc, 'Luồng thay thế và ngoại lệ'); label.runs[0].bold = True; label.paragraph_format.keep_with_next = True
        for i, alt in enumerate(uc['alts'], 1): p(doc, f"{uc['id']}-E{i}: {alt}")
        p(doc, f"Hậu điều kiện: {uc['post']}")
        p(doc, f"Quy tắc và nghiệm thu: {uc['rules']}")
    section_content(doc, DATA['usecase_back'])
    path = OUT / 'HUCEScholar_Use_Case_Specification.docx'; doc.save(path)
    return path


if __name__ == '__main__':
    paths = [proposal(), usecases()]
    result = {'outputs': [str(x) for x in paths], 'usecase_count': len(DATA['usecases']), 'functional_requirement_count': len(DATA['requirements'])}
    (QA / 'build_report.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))
