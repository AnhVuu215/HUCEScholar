"""Generate SQL, editable Mermaid, illustrated Word report and dictionary."""
from pathlib import Path
import json
import re
from PIL import Image, ImageDraw, ImageFont
from docx.shared import Cm, Pt
from build_hucescholar_docs import new_doc, p, heading, table, hyperlink, field
from hucescholar_data_model import TABLES, LINKS, INDEXES

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/database'
FIG = OUT / 'diagrams'
QA = ROOT / 'docs/_qa'
for folder in [OUT, FIG, QA]: folder.mkdir(parents=True, exist_ok=True)
BY_NAME = {t['name']: t for t in TABLES}
assert len(BY_NAME) == len(TABLES) == 23
for link in LINKS:
    assert link['child'] in BY_NAME and link['parent'] in BY_NAME
    assert link['column'] in {c['name'] for c in BY_NAME[link['child']]['columns']}


def sql_literal(value): return "'" + value.replace("'", "''") + "'"


def sql_schema():
    result = ['-- HUCEScholar reference schema 1.0 / 2026-10-10',
              '-- Target: PostgreSQL. Run only against a separate empty database.',
              '-- No DROP, seed data, database credentials or external file downloads.',
              'BEGIN;', 'CREATE SCHEMA hucescholar;',
              'REVOKE ALL ON SCHEMA hucescholar FROM PUBLIC;',
              'SET LOCAL search_path = hucescholar, pg_catalog;']
    for t in TABLES:
        rows = [f"  {c['name']} {c['sql']}" for c in t['columns']]
        rows += ['  ' + c for c in t['constraints']]
        result.append(f"CREATE TABLE {t['name']} (\n" + ',\n'.join(rows) + '\n);')
        result.append(f"COMMENT ON TABLE {t['name']} IS {sql_literal(t['purpose'])};")
        for c in t['columns']:
            result.append(f"COMMENT ON COLUMN {t['name']}.{c['name']} IS {sql_literal(c['meaning'])};")
    result += [
        'ALTER TABLE documents ADD CONSTRAINT fk_current_version_same_document FOREIGN KEY (id, current_version_id) REFERENCES document_versions(document_id, id) ON DELETE RESTRICT DEFERRABLE INITIALLY DEFERRED;',
        *INDEXES,
        (Path(__file__).parent / 'hucescholar_schema_guards.sql').read_text(encoding='utf-8'),
        'REVOKE ALL ON ALL TABLES IN SCHEMA hucescholar FROM PUBLIC;',
        'REVOKE ALL ON ALL SEQUENCES IN SCHEMA hucescholar FROM PUBLIC;',
        'REVOKE ALL ON ALL FUNCTIONS IN SCHEMA hucescholar FROM PUBLIC;',
        'COMMIT;'
    ]
    (OUT / 'HUCEScholar_schema_postgresql.sql').write_text('\n\n'.join(result) + '\n', encoding='utf-8')


def keys(t, c):
    marks = []
    if 'PRIMARY KEY' in c['sql'] or any(x.startswith('PRIMARY KEY') and c['name'] in re.findall(r'\w+', x)[2:] for x in t['constraints']): marks.append('PK')
    if any(l['child'] == t['name'] and l['column'] == c['name'] for l in LINKS): marks.append('FK')
    return ','.join(marks)


def typename(c):
    return re.match(r'^[a-z]+(?:\([\d,]+\))?', c['sql']).group(0)


def nullable(t, c):
    return 'Không' if 'NOT NULL' in c['sql'] or 'PK' in keys(t, c) else 'Có'


VIEWS = [
    ('01_logical', 'ERD tổng quan học liệu và học tập',
     ['users', 'courses', 'documents', 'files', 'document_versions', 'publication_requests', 'annotations', 'books'],
     [('users','documents','1','0..N'),('documents','document_versions','1','0..N'),('files','document_versions','1','0..N'),('document_versions','publication_requests','1','0..N'),('document_versions','annotations','1','0..N'),('users','annotations','1','0..N'),('courses','document_versions','0..N','1..N'),('courses','books','0..N','0..N')]),
    ('02_identity_catalog', 'ERD tài khoản và danh mục',
     ['faculties','programs','courses','program_courses','users','lecturer_assignments','auth_sessions','password_reset_tokens'],
     [('faculties','programs','1','0..N'),('faculties','courses','1','0..N'),('programs','program_courses','1','0..N'),('courses','program_courses','1','0..N'),('users','lecturer_assignments','1','0..N'),('courses','lecturer_assignments','1','0..N'),('users','auth_sessions','1','0..N'),('users','password_reset_tokens','1','0..N')]),
    ('03_content_review', 'ERD phiên bản và kiểm duyệt',
     ['documents','files','document_versions','publication_requests','version_courses','ownership_transfers','courses','visibility_events'],
     [('documents','document_versions','1','0..N'),('files','document_versions','1','0..N'),('document_versions','publication_requests','1','0..N'),('document_versions','version_courses','1','0..N'),('courses','version_courses','1','0..N'),('documents','ownership_transfers','1','0..N'),('documents','visibility_events','1','0..N')]),
    ('04_personal_library', 'ERD học tập cá nhân và sách',
     ['users','documents','annotations','document_versions','reading_progress','saved_documents','books','book_courses'],
     [('users','annotations','1','0..N'),('document_versions','annotations','1','0..N'),('users','reading_progress','1','0..N'),('document_versions','reading_progress','1','0..N'),('users','saved_documents','1','0..N'),('documents','saved_documents','1','0..N'),('books','book_courses','1','0..N'),('documents','books','0..1','0..N')]),
    ('05_reports_audit', 'ERD báo cáo và lịch sử xử lý',
     ['documents','books','document_versions','reports','visibility_events','report_actions','users','audit_events'],
     [('documents','reports','0..1','0..N'),('books','reports','0..1','0..N'),('document_versions','reports','0..1','0..N'),('reports','report_actions','1','0..N'),('visibility_events','report_actions','0..1','0..N'),('users','reports','1','0..N'),('users','audit_events','0..1','0..N')])
]


def mermaid():
    lines = ['erDiagram']
    for t in TABLES:
        lines.append('    ' + t['name'] + ' {')
        for c in t['columns']:
            typ = re.sub(r'\(.*', '', typename(c))
            mark = keys(t, c)
            lines.append(f"        {typ} {c['name']}" + (' ' + mark if mark else ''))
        lines.append('    }')
    for l in LINKS:
        left = '||' if l['required'] else '|o'
        right = 'o|' if l.get('child_max') == 1 else 'o{'
        child = BY_NAME[l['child']]
        column = next(c for c in child['columns'] if c['name'] == l['column'])
        # Identifying only when the referencing key participates in the child's PK.
        connector = '--' if 'PK' in keys(child, column) else '..'
        lines.append(f"    {l['parent']} {left}{connector}{right} {l['child']} : {l['column']}")
    (OUT / 'HUCEScholar_ERD_physical.mmd').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    logical = ['erDiagram']
    for a,b,c,d in VIEWS[0][3]:
        symbol = {'1':'||','0..1':'|o','0..N':'}o','1..N':'}|'}[c]
        rsymbol = {'1':'||','0..1':'o|','0..N':'o{','1..N':'|{'}[d]
        label = 'via_version_courses' if a == 'courses' and b == 'document_versions' else 'via_book_courses' if a == 'courses' and b == 'books' else 'relates_to'
        logical.append(f'    {a} {symbol}..{rsymbol} {b} : {label}')
    (OUT / 'HUCEScholar_ERD_logical.mmd').write_text('\n'.join(logical) + '\n', encoding='utf-8')


def draw_view(view):
    name, title, nodes, edges = view
    im = Image.new('RGB', (1800, 1580), 'white'); d = ImageDraw.Draw(im)
    titlefont = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 35)
    font = ImageFont.truetype('C:/Windows/Fonts/consola.ttf', 30)
    small = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 27)
    boxes = {}
    for i,n in enumerate(nodes):
        x = 45 if i%2 == 0 else 1055; y = 65 + (i//2)*365
        boxes[n] = (x,y,x+700,y+260)
    used = {n: 0 for n in nodes}
    for idx,(a,b,ca,cb) in enumerate(edges):
        ba,bb = boxes[a],boxes[b]
        ya = ba[1]+125+used[a]*25; yb = bb[1]+125+used[b]*25
        used[a]+=1; used[b]+=1
        xa = ba[2] if ba[0]<800 else ba[0]
        xb = bb[2] if bb[0]<800 else bb[0]
        lane = 790+idx*27
        d.line([(xa,ya),(lane,ya),(lane,yb),(xb,yb)],fill='#607386',width=3)
        for x,y,card,left in [(xa,ya,ca,ba[0]<800),(xb,yb,cb,bb[0]<800)]:
            d.text((x+(8 if left else -8),y-9),card,font=small,fill='black',anchor='lb' if left else 'rb',stroke_width=3,stroke_fill='white')
    for n,box in boxes.items():
        d.rectangle(box,fill='white',outline='#233B53',width=3)
        d.rectangle((box[0],box[1],box[2],box[1]+64),fill='#233B53')
        d.text((box[0]+17,box[1]+14),n,font=titlefont,fill='white')
        t = BY_NAME[n]
        fields = [(c['name'],keys(t,c)) for c in t['columns'] if keys(t,c)]
        if len(fields) < 4: fields += [(c['name'],'') for c in t['columns'] if not keys(t,c)][:4-len(fields)]
        for i,(col,mark) in enumerate(fields[:4]):
            d.text((box[0]+17,box[1]+81+i*37),(mark+' ' if mark else '')+col,font=font,fill='black')
    d.text((45,1537),'PK khóa chính   FK khóa ngoại   0..N không hoặc nhiều   1..N một hoặc nhiều',font=small,fill='black')
    target = FIG / (name+'.png'); im.save(target)
    return target


DECISIONS = [
    ('Tách tài liệu và phiên bản', 'documents giữ định danh, manager và visibility. document_versions giữ file, môn, metadata và quyền của snapshot. Ghi chú luôn tham chiếu version_id, không tham chiếu current_version_id.'),
    ('Môn thuộc phiên bản', 'Quan hệ học phần nằm ở version_courses. Bản nháp phải có ít nhất một môn trước gửi duyệt. SQL cho phép nháp chưa chọn môn; backend kiểm tra điều kiện gửi. Bàn giao kiểm tra người nhận được phân công hợp của tất cả môn trong các phiên bản thuộc tài liệu.'),
    ('Tác giả và chủ quản lý', 'documents.created_by và versions.created_by lưu người tạo gốc. authors là nhãn tác giả học thuật; manager_id là người quản lý hiện tại. publication_requests.submitted_by giữ người gửi thực sự của từng lần, không đổi sau bàn giao.'),
    ('Một role chính và một bản đang làm', 'Lựa chọn thiết kế MVP: users.role có một giá trị; một tài liệu chỉ có một bản DRAFT, SUBMITTED hoặc CHANGES_REQUIRED tại một thời điểm. Không bổ sung role kiểm duyệt riêng. Có thể thay đổi mô hình này nếu sau này cần kiêm nhiệm hoặc nhiều nhánh biên tập.'),
    ('Tài liệu và sách nguồn ngoài', 'books chứa metadata, nguồn và quyền dùng metadata. native_document_id là tùy chọn, không chứng minh quyền phục vụ toàn văn. Khi mở PDF vẫn kiểm tra quyền của document/version. Không có dữ liệu tồn kho, mượn trả hoặc API trường trong MVP.'),
    ('Xóa và lưu lịch sử', 'Các FK dùng RESTRICT để tránh xóa dây chuyền file, phiên bản và dữ liệu riêng tư. Ngừng tài khoản, danh mục hoặc ẩn nội dung bằng trạng thái. Chính sách xóa dữ liệu thật và thời hạn lưu cần thống nhất trước vận hành với người dùng thật.')
]
RULES = [
    ('Database', 'PK, NOT NULL, CHECK giá trị trạng thái; FK và quan hệ phiên bản thuộc tài liệu; unique email, số phiên bản, mục lưu và tiến độ.'),
    ('Database', 'Partial unique: một phân công chưa thu hồi mỗi giảng viên/môn, một bản đang làm mỗi tài liệu và một lần duyệt PENDING mỗi phiên bản.'),
    ('Database', 'Trigger đóng băng bản đã APPROVED/REJECTED và payload đang SUBMITTED; bảo vệ các môn của snapshot, danh tính file và tác giả gốc. Con trỏ hiện hành phải cùng tài liệu và APPROVED ở cuối giao dịch.'),
    ('Backend', 'Kiểm tra ACTIVE, verified_at, role hiện tại, sở hữu và tất cả phân công ở mỗi thao tác. FK user không chứng minh người đó là giảng viên/admin đang hoạt động.'),
    ('Backend', 'Kiểm tra quyền nguồn, SAFE, thời hạn quyền, VISIBLE và APPROVED tại mỗi lần cấp file. Preview riêng theo quyền quản lý chỉ khi file SAFE. Kiểm tra MIME thật và bảo vệ object storage.'),
    ('Backend', 'page_number không vượt page_count; highlight có lớp chữ ở trang đó, quads đúng cấu trúc và nằm trong 0..1. SQL chỉ kiểm tra dạng JSON cơ bản và trang dương.'),
    ('Backend', 'Cập nhật revision theo WHERE id = ? AND revision = ?; không ghi đè xung đột. Giữ admin hoạt động cuối cùng bằng khóa giao dịch chung trước thay role/trạng thái.'),
    ('Backend', 'Chuyển trạng thái hợp lệ, lịch sử append only, hạn mức upload, chống spam, token một lần/hết hạn và thu hồi session. Không cho role ADMIN đọc ghi chú của người khác.'),
    ('Kho file', 'object_key là khóa riêng tư bất biến; không ghi đè bytes tại cùng khóa. Hash và kiểm tra malware thực tế không được suy ra từ metadata SQL.')
]
TRANSACTIONS = [
    ('Gửi duyệt', 'Khóa documents rồi phiên bản; kiểm tra quyền và revision. Với bản bị yêu cầu sửa, bắt đầu chỉnh sửa đưa về DRAFT. Kiểm tra SAFE, nguồn/quyền và ít nhất một môn; tăng revision khi chuyển SUBMITTED rồi tạo publication_requests PENDING với đúng submitted_revision.'),
    ('Duyệt công bố', 'Khóa documents, phiên bản rồi request theo thứ tự cố định. So sánh request PENDING, phiên bản SUBMITTED và submitted_revision. Kiểm tra lại SAFE, môn/phân công và quyền; kết thúc request, đổi phiên bản APPROVED và current_version_id trong cùng giao dịch. Giữ HIDDEN nếu tài liệu đang ẩn.'),
    ('Rút và gửi lại', 'Cùng khóa với duyệt. Kết thúc request bằng WITHDRAWN, ghi người/lý do; đưa phiên bản về DRAFT và tăng revision. Gửi lại tạo request mới, giữ nguyên lần gửi trước. Khi admin yêu cầu sửa, request kết thúc CHANGES_REQUIRED; phiên bản cũng mang trạng thái đó tới khi bắt đầu sửa.'),
    ('Bàn giao giảng viên', 'Khóa document trước các phiên bản/request. Chỉ LECTURER_MATERIAL; kiểm tra revision và người nhận ACTIVE, LECTURER, có mọi môn liên quan. Đổi manager; kết thúc request PENDING bằng WITHDRAWN với lý do bàn giao; bản SUBMITTED về DRAFT, tăng revision. Ghi ownership_transfers và audit trong cùng giao dịch; giữ file, tác giả, ghi chú và bản công bố.'),
    ('Ẩn hoặc khôi phục', 'Khóa document, kiểm tra revision và lý do. Cập nhật visibility, visibility_events, audit cùng giao dịch. Khôi phục cần bản hiện hành APPROVED và SAFE, quyền hợp lệ; kiểm tra riêng bản cũ khi được yêu cầu đọc. Dừng lượt cấp file mới, không hứa thu hồi file đã tải.'),
    ('Ghi chú và tiến độ', 'Lọc owner bằng user trong phiên máy chủ, kiểm tra quyền đọc phiên bản và page_count, cập nhật theo revision. Anchor dùng tọa độ chuẩn hóa 0..1 trên trang PDF gốc chưa xoay; không lấy tọa độ màn hình. Nội dung ghi chú không vào audit hoặc thông báo cho giảng viên.')
]
COVERAGE = [
    ('UC01–UC04', 'users, auth_sessions, password_reset_tokens', 'Đăng nhập, xác minh, khóa/role và khôi phục mật khẩu'),
    ('UC05–UC06', 'faculties, programs, courses, program_courses, lecturer_assignments, ownership_transfers', 'Danh mục, phân công và bàn giao'),
    ('UC07–UC10', 'documents, document_versions, version_courses, files, annotations, reading_progress', 'Tìm, đọc, ghi chú và tải'),
    ('UC11–UC12', 'saved_documents, annotations, reading_progress, books, book_courses', 'Thư viện cá nhân và nguồn sách'),
    ('UC13–UC18', 'documents, document_versions, files, version_courses, publication_requests', 'Đóng góp, biên soạn, phiên bản và công bố'),
    ('UC19–UC21', 'visibility_events, reports, report_actions, audit_events', 'Ẩn/khôi phục, báo cáo và xử lý'),
    ('UC22', 'books, book_courses, documents', 'Biên tập danh mục sách và liên kết PDF được phép')
]


def dictionary():
    rows = ['# HUCEScholar Data Dictionary', '', 'Phiên bản thiết kế 1.0 ngày 10/10/2026. Cơ sở nghiệp vụ 1.1.', '',
            f'{len(TABLES)} bảng; {sum(len(t["columns"]) for t in TABLES)} trường. SQL là phương án tham chiếu PostgreSQL, chưa phải database triển khai.', '',
            '## Quyết định thiết kế', '']
    for name, text in DECISIONS: rows += ['### '+name, '', text, '']
    for t in TABLES:
        rows += ['## '+t['name'], '', t['purpose'], '', '| Trường | Kiểu | Khóa | Null | Ý nghĩa |', '|---|---|---|---|---|']
        for c in t['columns']:
            rows.append(f"| {c['name']} | {typename(c)} | {keys(t,c) or '—'} | {nullable(t,c)} | {c['meaning']} |")
        if t['constraints']: rows += ['', 'Ràng buộc: ' + '; '.join('`'+c+'`' for c in t['constraints']), '']
        for l in LINKS:
            if l['child'] == t['name']: rows += [f"- `{l['column']}` → `{l['parent']}.id`" + ('; FK thực tế dùng cặp document_id và version_id.' if l['column'] in ['current_version_id','version_id'] and t['name'] in ['documents','reports'] else '')]
        rows.append('')
    rows += ['## Phân chia trách nhiệm', '']
    for who,text in RULES: rows += [f'- **{who}:** {text}']
    rows += ['', '## Giao dịch trọng yếu', '']
    for name,text in TRANSACTIONS: rows += ['### '+name, '', text, '']
    rows += ['## Chỉ mục', '', '```sql', *INDEXES, '```', '',
             'GIN chỉ tìm metadata chuẩn hóa, không phải toàn văn PDF. Backend tạo search_text và chuẩn hóa truy vấn cùng quy tắc tiếng Việt. Chỉ mục không tự chứng minh đáp ứng mục tiêu hiệu năng; cần đo trên dữ liệu và môi trường thực.', '',
             '## Truy vết use case', '', '| Use case | Bảng | Mục đích |', '|---|---|---|']
    rows += ['| '+' | '.join(r)+' |' for r in COVERAGE]
    rows += ['', '## Giới hạn kiểm chứng', '', 'Kết quả chạy schema và ca kiểm tra database tạm được ghi ở docs/_qa/data_design_validation.json. Chưa có backend, RLS, kho PDF hoặc migration lên database triển khai. Fixture rollback không phải seed data lâu dài.', '',
             '## Tài liệu kỹ thuật', '',
             '- PostgreSQL Constraints: https://www.postgresql.org/docs/18/ddl-constraints.html',
             '- PostgreSQL CREATE TRIGGER: https://www.postgresql.org/docs/18/sql-createtrigger.html', '']
    (OUT / 'HUCEScholar_Data_Dictionary.md').write_text('\n'.join(rows), encoding='utf-8')


def word_report():
    doc = new_doc('ERD and Data Design')
    # New design version/date; preserve styles of the existing project documents.
    foot = doc.sections[0].footer.paragraphs[0]
    foot.clear(); foot.alignment = 2
    foot.add_run('Phiên bản 1.0  •  10/10/2026  |  Trang ').font.size = Pt(9)
    field(foot, ' PAGE ')
    doc.styles['Normal'].font.size = Pt(11.5)
    p(doc, 'HUCEScholar Data Design', 'Title')
    p(doc, 'ERD và từ điển dữ liệu cho MVP', 'Subtitle')
    p(doc, 'Phiên bản 1.0  |  Ngày 10/10/2026  |  Cơ sở nghiệp vụ 1.1')
    p(doc, 'Thiết kế tổ chức dữ liệu cho nền tảng giáo trình điện tử HUCEScholar với ba vai trò ADMIN, LECTURER và STUDENT. Mục tiêu là giữ đúng nguồn học liệu, lịch sử công bố và dữ liệu học tập riêng tư khi thay PDF hoặc bàn giao giảng viên.')
    p(doc, f'Mô hình gồm {len(TABLES)} bảng và {sum(len(t["columns"]) for t in TABLES)} trường. Tài liệu cung cấp ERD tổng quan, các góc nhìn vật lý, từ điển dữ liệu, trách nhiệm kiểm tra và giao dịch trọng yếu. Schema SQL PostgreSQL và Mermaid đi kèm để triển khai đối chiếu.')
    p(doc, 'Phạm vi giữ nguyên MVP đã chốt. Chưa có lớp học/đăng ký môn, OCR, AI, ghi chú cộng tác, mượn sách hoặc đồng bộ API trường. PostgreSQL là lựa chọn tham chiếu cho thiết kế; quyền dữ liệu trường và toàn bộ stack ứng dụng chưa được xác nhận.')
    heading(doc, '1 Quyết định thiết kế')
    for name,text in DECISIONS:
        heading(doc, name, 2); p(doc,text)
    doc.add_page_break()
    heading(doc, '2 Sơ đồ quan hệ dữ liệu')
    p(doc, 'Hình tổng quan mô tả quan hệ nghiệp vụ. Các hình vật lý tiếp theo thể hiện bảng và khóa chính/ngoại chính của từng nhóm; để đọc rõ, không vẽ lặp tất cả FK tới người thực hiện lịch sử. File Mermaid vật lý và SQL liệt kê đầy đủ các FK. Bảng xuất hiện ở nhiều hình là cùng một bảng.')
    for i,view in enumerate(VIEWS):
        if i: doc.add_page_break()
        heading(doc, f'2 {i+1} '+view[1], 2)
        par=p(doc,''); run=par.add_run(); run.add_picture(str(draw_view(view)),width=Cm(16.7))
        run._r.xpath('.//wp:docPr')[0].set('descr', view[1]+'; các bảng '+', '.join(view[2]))
        p(doc, f'Hình {i+1} '+view[1], 'Caption')
        if i == 0: p(doc, 'Quan hệ môn với phiên bản và sách là nhiều–nhiều, được triển khai bằng version_courses và book_courses. Ở tầng lưu trữ, phiên bản nháp có thể chưa chọn môn; trước gửi duyệt cần ít nhất một môn.')
        if i == 2: p(doc, 'documents.current_version_id tham chiếu đúng phiên bản của chính tài liệu và chỉ trỏ bản APPROVED ở cuối giao dịch. Một phiên bản có nhiều lần gửi duyệt theo thời gian, nhưng tối đa một lần PENDING.')
        if i == 4: p(doc, 'reports chọn đúng một đối tượng: document hoặc book. version_id là tùy chọn và phải thuộc document_id. audit_events.entity_id là tham chiếu đa hình không có FK; backend xác thực và chỉ ghi trường đã cho phép.')
    doc.add_page_break(); heading(doc, '3 Từ điển dữ liệu')
    p(doc, 'PK là khóa chính, FK là khóa ngoại; Null Có nghĩa là có thể để trống. Các ID dùng bigint identity; thời gian dùng timestamptz. Mọi bảng nối có khóa chính kép. Không lưu bytes PDF trong database. Danh sách DEFAULT và CHECK đầy đủ nằm trong SQL đi kèm.')
    group = None
    for i,t in enumerate(TABLES,1):
        if t['group'] != group:
            if group is not None: doc.add_page_break()
            group=t['group']; heading(doc, group, 2)
        heading(doc, f"3 {i} {t['name']}", 3)
        p(doc,t['purpose']).paragraph_format.keep_with_next = True
        headers=['Trường và khóa','Kiểu','Null','Ý nghĩa']
        rows=[[c['name'] + (' ['+keys(t,c)+']' if keys(t,c) else ''), typename(c), nullable(t,c), c['meaning']] for c in t['columns']]
        built = table(doc,headers,rows,[4.9,2.5,1.6,7.8])
        for row in built.rows[:3]:
            for cell in row.cells:
                for par in cell.paragraphs: par.paragraph_format.keep_with_next = True
        # Explain compound rules in readable prose, rather than squeezing SQL into cells.
        if t['name'] == 'documents': p(doc, 'FK kép (id, current_version_id) buộc phiên bản thuộc đúng tài liệu. Trigger trì hoãn kiểm tra bản hiện hành APPROVED khi kết thúc giao dịch.')
        if t['name'] == 'document_versions': p(doc, 'Duy nhất (document_id, version_no); tối đa một bản đang làm DRAFT/SUBMITTED/CHANGES_REQUIRED. Snapshot đã duyệt hoặc bị từ chối không sửa; tạo phiên bản mới.')
        if t['name'] == 'reports': p(doc, 'CHECK buộc đúng một trong document_id và book_id. FK kép (document_id, version_id) ngăn báo cáo nhầm bản của tài liệu khác.')
    doc.add_page_break(); heading(doc,'4 Ràng buộc và trách nhiệm kiểm tra')
    for who,text in RULES:
        par=p(doc,''); par.add_run(who+': ').bold=True; par.add_run(text)
    heading(doc,'4 1 Chỉ mục và tìm kiếm',2)
    p(doc,'Chỉ mục ưu tiên hàng đợi duyệt/báo cáo, phiên bản theo tài liệu, môn liên quan, ghi chú theo người–phiên bản–trang, tiến độ và mục lưu. Partial unique bảo vệ bản đang làm, lần duyệt PENDING và phân công chưa thu hồi. FK không tự tạo chỉ mục ở phía tham chiếu; các FK truy vấn chính có chỉ mục riêng trong SQL.')
    p(doc,'GIN tìm metadata search_text được chuẩn hóa bởi backend, không tìm nội dung PDF. Cần chuẩn hóa truy vấn cùng quy tắc tiếng Việt và đo mục tiêu tìm kiếm trên môi trường thật; chưa có kết quả benchmark.')
    heading(doc,'4 2 Phạm vi bảo mật dữ liệu',2)
    p(doc,'Schema thu hồi quyền PUBLIC nhưng chưa tạo role kết nối cho backend hoặc bật Row Level Security. Không cấp database credential cho người dùng cuối. API phải lọc owner và kiểm tra quyền ở máy chủ; admin ứng dụng không có quyền đọc ghi chú người khác. Session/reset token chỉ lưu hash. Audit không chứa payload ghi chú hoặc bí mật.')
    doc.add_page_break(); heading(doc,'5 Giao dịch trọng yếu')
    for name,text in TRANSACTIONS: heading(doc,name,2); p(doc,text)
    heading(doc,'5 1 Thứ tự khóa và chống cập nhật đồng thời',2)
    p(doc,'Mọi luồng thay đổi cùng một học liệu khóa document trước, sau đó version rồi request theo ID tăng dần. Tạo số phiên bản dưới khóa document; không dùng MAX(version_no)+1 ngoài giao dịch. Duyệt, rút và bàn giao phải dùng cùng thứ tự khóa. Thu hồi phân công và duyệt cần đồng bộ khóa dữ liệu phân công để không công bố từ một quyền vừa hết hiệu lực.')
    doc.add_page_break(); heading(doc,'6 Truy vết và kiểm chứng')
    table(doc,['Use case','Bảng dữ liệu','Mục đích'],COVERAGE,[2.5,7.8,6.5])
    heading(doc,'6 1 Phạm vi kiểm chứng',2)
    p(doc,'Schema được kiểm tra bằng bộ script database tạm độc lập; kết quả chi tiết ghi trong data_design_validation.json. Các ca kiểm tra tập trung vào FK, trùng dữ liệu, đóng băng phiên bản, con trỏ hiện hành, báo cáo đúng đối tượng và giữ liên kết ghi chú. Dữ liệu fixture rollback, không phải kho học liệu thật hoặc seed tồn tại lâu dài.')
    p(doc,'Kiểm tra schema không chứng minh backend phân quyền đúng, giao dịch đồng thời hoạt động, PDF đọc/ghi chú được, dữ liệu trường đã nhập hoặc migration lên database triển khai đã thành công. Những phần này cần kiểm thử khi xây ứng dụng.')
    heading(doc,'6 2 Công việc tiếp theo',2)
    p(doc,'Dựa vào ERD để thiết kế wireframe ba vai trò: tìm–đọc–ghi chú, tạo–gửi–duyệt, bàn giao và xử lý báo cáo. Thử kỹ thuật reader trên PDF có chữ, scan và trang xoay. Sau đó chốt stack, ánh xạ session của framework, viết API và migration theo mô hình đã được rà soát.')
    heading(doc,'7 Tài liệu tham chiếu')
    p(doc,'Cơ sở nghiệp vụ: HUCEScholar Project Proposal, Use Case Specification và Business Rules and Publication Workflow phiên bản 1.1. Các tài liệu kỹ thuật dưới đây dùng để đối chiếu FK, CHECK, partial unique và trigger; không xác nhận quyền dữ liệu của trường.')
    hyperlink(doc,'PostgreSQL Constraints','https://www.postgresql.org/docs/18/ddl-constraints.html')
    hyperlink(doc,'PostgreSQL CREATE TRIGGER','https://www.postgresql.org/docs/18/sql-createtrigger.html')
    target=ROOT/'docs/deliverables/HUCEScholar_ERD_and_Data_Design.docx'; doc.save(target)
    return target


if __name__ == '__main__':
    sql_schema(); mermaid(); dictionary(); report=word_report()
    (OUT/'HUCEScholar_model.json').write_text(json.dumps({'tables':TABLES,'links':LINKS,'indexes':INDEXES},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'tables':len(TABLES),'fields':sum(len(t['columns']) for t in TABLES),'relationships':len(LINKS),'report':str(report)},ensure_ascii=False))
