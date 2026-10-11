from pathlib import Path
import json
import re
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from PIL import Image, ImageDraw, ImageFont

root = Path(__file__).resolve().parents[2]
docs = root / 'docs'
qa = docs / '_qa'
data = json.loads((Path(__file__).parent / 'hucescholar_content.json').read_text(encoding='utf-8'))
ids = [x['id'] for x in data['usecases']]
frids = [x[0] for x in data['requirements']]
assert ids == [f'UC{i:02}' for i in range(1,23)]
assert frids == [f'FR{i:02}' for i in range(1,13)]
required = ['id','name','goal','actors','fr','pre','trigger','steps','alts','post','rules']
for uc in data['usecases']:
    assert all(uc.get(k) for k in required), uc['id']
    assert all(fr in frids for fr in uc['fr']), uc['id']
for fr in frids:
    assert any(fr in uc['fr'] for uc in data['usecases']), fr

report = {'schema_checks': 'passed', 'usecases': 22, 'functional_requirements': 12, 'documents': []}
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 24)
for stem, image_dir in [('HUCEScholar_Project_Proposal','proposal_pages_v4'),('HUCEScholar_Use_Case_Specification','usecase_pages_v4'),('HUCEScholar_Business_Rules_and_Publication_Workflow','business_pages_v4')]:
    document_path = docs / 'deliverables' / f'{stem}.docx'
    doc = Document(document_path)
    text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert 'HUCEReader' not in text
    assert not re.search(r'\b(TODO|TBD|FIXME)\b', text)
    assert '\ufffd' not in text
    if 'Business_Rules' not in stem:
        assert all(fr in text for fr in frids)
    assert 'Phiên bản 1.1' in text
    assert 'Theo quyền quản trị' not in text
    assert 'Chưa có quyết định của người đề xuất' not in text
    for aid in [f'A{i:02}' for i in range(1,7)]: assert aid in text
    if 'Use_Case' in stem:
        headings = [p.text for p in doc.paragraphs if p.style.name == 'Heading 2' and re.match(r'^UC\d{2} ', p.text)]
        assert len(headings) == 22
        assert len(doc.inline_shapes) == 3
        assert all(f'BR{i:02}' in text for i in range(1,18))
    if 'Business_Rules' in stem:
        assert len(doc.inline_shapes) == 1
        assert all(word in text for word in ['BR16', 'BR17', 'CHANGES_REQUIRED', 'SUBMITTED', 'APPROVED', 'REJECTED'])
    for t in doc.tables:
        assert t.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
        assert t._tbl.tblPr.find(qn('w:tblBorders')) is not None
    assert not doc.styles['Title'].element.xpath('./w:pPr/w:pBdr')
    with ZipFile(document_path) as z:
        assert not any('vbaProject' in name for name in z.namelist())
    pdf = PdfReader(qa / f'{stem}.pdf')
    pages = list((qa / image_dir).glob('page-*.png'))
    pages.sort(key=lambda p: int(re.search(r'(\d+)\.png$', p.name).group(1)))
    assert len(pages) == len(pdf.pages)
    page_summaries = []
    for i, page in enumerate(pdf.pages,1):
        txt = page.extract_text() or ''
        assert '\ufffd' not in txt
        page_summaries.append({'page':i,'characters':len(txt),'first_lines':txt.splitlines()[:5],'last_lines':txt.splitlines()[-4:]})
    # Contact sheets are only an overview; native page PNGs are retained for inspection.
    for start in range(0,len(pages),4):
        sheet = Image.new('RGB',(1040,1540),'#E5E5E5')
        draw = ImageDraw.Draw(sheet)
        for offset,image_path in enumerate(pages[start:start+4]):
            img = Image.open(image_path); img.thumbnail((500,710))
            x = 10+(offset%2)*520; y=40+(offset//2)*770
            draw.text((x,y-30),f'{image_dir} {start+offset+1}',font=font,fill='black')
            sheet.paste(img,(x,y))
        sheet.save(qa / f'{image_dir}_contact_{start//4+1}.png')
    report['documents'].append({'file':str(document_path),'bytes':document_path.stat().st_size,'pages':len(pdf.pages),'tables':len(doc.tables),'images':len(doc.inline_shapes),'page_summaries':page_summaries})
(qa / 'structural_qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'schema_checks':report['schema_checks'],'documents':[{k:v for k,v in d.items() if k!='page_summaries'} for d in report['documents']]},ensure_ascii=False,indent=2))
for d in report['documents']:
    print('\n'+Path(d['file']).name)
    for page in d['page_summaries']:
        print(f"p{page['page']:02} chars={page['characters']} start="+' / '.join(page['first_lines'][1:4]))
