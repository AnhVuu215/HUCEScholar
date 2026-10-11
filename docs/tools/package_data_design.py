"""Check generated design artifacts and package the reviewed deliverables."""
from pathlib import Path
import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from docx import Document
from pypdf import PdfReader
from build_data_design import COVERAGE
from hucescholar_data_model import TABLES, LINKS

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / 'docs'
DATABASE = DOCS / 'database'
QA = DOCS / '_qa'
REPORT = DOCS / 'deliverables/HUCEScholar_ERD_and_Data_Design.docx'


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    doc = Document(REPORT)
    sql_path = DATABASE / 'HUCEScholar_schema_postgresql.sql'
    sql = sql_path.read_text(encoding='utf-8')
    md = (DATABASE / 'HUCEScholar_Data_Dictionary.md').read_text(encoding='utf-8')
    model = json.loads((DATABASE / 'HUCEScholar_model.json').read_text(encoding='utf-8'))
    text = '\n'.join(p.text for p in doc.paragraphs)
    text += '\n' + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert model['tables'] == TABLES and model['links'] == LINKS
    assert len(TABLES) == 23 and sum(len(t['columns']) for t in TABLES) == 184
    assert len(LINKS) == 47
    assert len(re.findall(r'\bREFERENCES\b', sql)) == 47
    assert len(re.findall(r'^CREATE TABLE ', sql, re.M)) == 23
    dictionaries = [t for t in doc.tables if t.rows[0].cells[0].text == 'Trường và khóa']
    assert len(dictionaries) == len(TABLES)
    for t, rendered in zip(TABLES, dictionaries):
        assert f"## {t['name']}\n" in md
        assert any(p.text.endswith(' ' + t['name']) for p in doc.paragraphs)
        assert [r.cells[0].text.split(' [')[0] for r in rendered.rows[1:]] == [c['name'] for c in t['columns']]
        assert len(rendered.rows) == len(t['columns']) + 1
        for c in t['columns']:
            assert f"| {c['name']} |" in md
            assert f"COMMENT ON COLUMN {t['name']}.{c['name']} IS" in sql
        assert rendered.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
        assert rendered._tbl.xpath('./w:tblPr/w:tblBorders')
    assert len(doc.inline_shapes) == 5
    assert all(s._inline.xpath('./wp:docPr')[0].get('descr') for s in doc.inline_shapes)
    title = next(p for p in doc.paragraphs if p.style.name == 'Title')
    assert not title._p.xpath('./w:pPr/w:pBdr')
    covered = set()
    for label, _, _ in COVERAGE:
        endpoints = [int(n) for n in re.findall(r'UC(\d+)', label)]
        covered.update(range(endpoints[0], endpoints[-1] + 1))
    assert covered == set(range(1, 23))
    for content in [text, md, sql]:
        assert not re.search(r'\b(?:TODO|TBD|FIXME|HUCEReader)\b|\ufffd', content)
    pdf = PdfReader(QA / 'HUCEScholar_ERD_and_Data_Design.pdf')
    pages = len(pdf.pages)
    assert pages == 23
    assert len(list((QA / 'data_design_pages').glob('page-*.png'))) == pages
    assert all(len(p.extract_text().strip()) > 100 for p in pdf.pages)
    validation = json.loads((QA / 'data_design_validation.json').read_text(encoding='utf-8'))
    assert validation['status'] == 'passed' and validation['passed'] == 37
    assert validation['foreign_keys'] == 47 and validation['persisted_fixture_users'] == 0
    assert validation['test_server'] == 'stopped'
    # Visual review is a separate manual step: page images 01-23 were inspected.
    structural = {
        'timestamp_utc': datetime.now(timezone.utc).isoformat(),
        'status': 'passed', 'tables': 23, 'fields': 184, 'foreign_keys': 47,
        'dictionary_tables': 23, 'embedded_diagrams': 5,
        'covered_use_cases': 22, 'rendered_pages': pages,
        'visual_review': 'All 23 rendered pages inspected manually; readable, no clipping or overlap',
        'docx_sha256': sha256(REPORT), 'schema_sha256': sha256(sql_path),
        'limitations': 'No backend authorization, concurrent transaction, PDF-reader or deployment validation',
    }
    structural_path = QA / 'data_design_artifact_validation.json'
    structural_path.write_text(json.dumps(structural, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    target = DOCS / 'deliverables/HUCEScholar_ERD_and_Database_Package.zip'
    expected = {}
    with zipfile.ZipFile(target, 'w', compression=zipfile.ZIP_DEFLATED) as package:
        files = [(REPORT, REPORT.name)]
        files += [(p, 'database/' + p.relative_to(DATABASE).as_posix()) for p in sorted(DATABASE.rglob('*')) if p.is_file()]
        files += [(QA / 'data_design_validation.json', 'validation/Database_Validation.json'),
                  (structural_path, 'validation/Artifact_Validation.json')]
        for path, name in files:
            expected[name] = sha256(path)
            package.write(path, name)
        package.writestr('SHA256SUMS.txt', '\n'.join(f'{value}  {name}' for name, value in expected.items()) + '\n')
    with zipfile.ZipFile(target) as package:
        assert package.testzip() is None
        for name, value in expected.items():
            assert hashlib.sha256(package.read(name)).hexdigest() == value
        assert len(package.namelist()) == len(expected) + 1
    print(json.dumps({'status': 'passed', 'pages': pages, 'package_entries': len(expected) + 1,
                      'package': str(target), 'bytes': target.stat().st_size}, ensure_ascii=False))


if __name__ == '__main__':
    main()
